"""Drafting the document a call actually asked for, editing it, exporting it.

Written after the owner read what the system produced for a POTRAZ research
call - a call wanting a 3.5-page, nine-section research proposal in Times New
Roman 12 - and got a one-paragraph cover letter. The letter was not a worse
proposal; it was the wrong artefact, and these tests pin the difference.
"""

import json

import httpx
from fastapi.testclient import TestClient

from conftest import sign_up
from opportunity_agent import api, application_spec, models_db, section_drafting
from opportunity_agent import db as db_module
from opportunity_agent.api import app
from opportunity_agent.application_spec import SectionSpec, SubmissionSpec

client = TestClient(app)


def setup_function():
    client.cookies.clear()
    sign_up(client)


def _client_with(handler):
    return httpx.Client(transport=httpx.MockTransport(handler))


def _reply(content):
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"message": {"role": "assistant", "content": content}})
    return handler


POTRAZ = """Call for Research Proposals. Proposals must not exceed 3.5 pages excluding the
cover page, in Times New Roman size 12, 1.5 spacing, submitted as PDF in English.
Structure: 1. Title of the Proposed Policy Research 2. Priority Research Area
3. Policy Problem Statement 4. Background and Justification 5. Research Objectives.
Submit to research.development@potraz.zw by 3 October 2026. Open to policy analysts,
researchers, academics."""


# --------------------------------------------------------------------------
# Reading the call
# --------------------------------------------------------------------------

def test_the_spec_reports_the_format_rules_the_call_states():
    answer = json.dumps({
        "document_kind": "policy research proposal",
        "sections": [{"title": "Policy Problem Statement", "guidance": "the challenge"}],
        "format_rules": {"max_pages": 3.5, "font": "Times New Roman", "font_size": 12,
                         "line_spacing": "1.5", "file_format": "PDF", "language": "English",
                         "other": ["AI-generated text below 10%"]},
        "submit_to": "research.development@potraz.zw",
        "deadline": "3 October 2026",
        "eligibility": ["policy analysts", "academics"],
    })

    spec = application_spec.extract_submission_spec(POTRAZ, client=_client_with(_reply(answer)))

    assert spec.format_rules.max_pages == 3.5
    assert spec.format_rules.font == "Times New Roman"
    assert spec.format_rules.font_size == 12
    assert spec.format_rules.file_format == "PDF"
    assert spec.submit_to == "research.development@potraz.zw"
    assert spec.deadline == "3 October 2026"
    # a rule that fits no field must still reach the owner
    assert any("10%" in o for o in spec.format_rules.other)


def test_sections_keep_the_calls_own_order_and_wording():
    answer = json.dumps({"sections": [
        {"title": "Title of the Proposed Policy Research"},
        {"title": "Priority Research Area"},
        {"title": "Policy Problem Statement"},
    ]})

    spec = application_spec.extract_submission_spec("x", client=_client_with(_reply(answer)))

    assert [s.title for s in spec.sections] == [
        "Title of the Proposed Policy Research",
        "Priority Research Area",
        "Policy Problem Statement",
    ]
    assert spec.is_structured


def test_a_call_with_no_readable_structure_yields_no_sections():
    """Better to fall back to the plain cover note than to invent a
    nine-section skeleton for a call that wanted two paragraphs."""
    spec = application_spec.extract_submission_spec(
        "Send us your CV.", client=_client_with(_reply(json.dumps({"sections": []})))
    )

    assert spec.sections == []
    assert not spec.is_structured


def test_an_unreadable_answer_does_not_raise():
    spec = application_spec.extract_submission_spec("x", client=_client_with(_reply("not json")))

    assert spec.sections == []


def test_a_bare_list_of_headings_is_still_accepted():
    """Some models answer with strings rather than objects - that is still
    the ordered structure that was asked for."""
    answer = json.dumps({"sections": ["Problem Statement", "Objectives"]})

    spec = application_spec.extract_submission_spec("x", client=_client_with(_reply(answer)))

    assert [s.title for s in spec.sections] == ["Problem Statement", "Objectives"]


# --------------------------------------------------------------------------
# Writing a section
# --------------------------------------------------------------------------

def test_the_section_prompt_offers_only_the_profiles_own_facts():
    """The anti-fabrication guarantee drafting.py had: a claim about the
    applicant must trace to something they entered."""
    class Profile:
        name = "Isaiah Chikeya"
        field = "computer systems engineering"
        certificates = ["BSc computer systems engineering"]
        work_history = ["AI Researcher HIT (Jan 2026 to date)"]

    facts = section_drafting.profile_facts(Profile())
    prompt = section_drafting.build_section_prompt(
        SectionSpec(title="Policy Problem Statement", guidance="the challenge"),
        SubmissionSpec(document_kind="research proposal"),
        type("O", (), {"title": "POTRAZ call", "summary": "about AI policy"})(),
        facts,
    )

    assert "Isaiah Chikeya" in prompt
    assert "AI Researcher HIT" in prompt
    assert "Do not invent qualifications" in prompt
    assert "Policy Problem Statement" in prompt


def test_a_section_that_cannot_be_written_is_empty_not_fatal():
    """Losing the whole package because section 6 of 9 timed out would be
    worse than a gap the owner can fill."""
    def boom(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("model down")

    body = section_drafting.draft_section(
        SectionSpec(title="Objectives"), SubmissionSpec(),
        type("O", (), {"title": "t", "summary": "s"})(), "- Name: X",
        client=_client_with(boom),
    )

    assert body == ""


# --------------------------------------------------------------------------
# The endpoints
# --------------------------------------------------------------------------

def _opportunity(profile_id):
    session = db_module.SessionLocal()
    row = models_db.StoredOpportunity(
        profile_id=profile_id, canonical_url="https://potraz.gov.zw/call",
        payload={"title": "Call for Research Proposals", "content": POTRAZ},
        match_status="eligible", match_score=10, match_reasons={}, stage="drafted",
    )
    session.add(row)
    session.commit()
    row_id = row.id
    session.close()
    return row_id


def _profile():
    p = client.post(
        "/profiles", json={"profile_type": "grant", "display_name": "Grants"}
    ).json()
    client.put(f"/profiles/{p['id']}", json={"fields": {"name": "Isaiah Chikeya"}})
    return p


def test_drafting_a_full_application_is_accepted_and_runs_in_the_background(monkeypatch):
    """It costs minutes of model time, so it must not block the request - and
    must not run unprompted for every discovered opportunity either."""
    profile = _profile()
    opp_id = _opportunity(profile["id"])
    scheduled = {}

    monkeypatch.setattr(api, "_full_draft_worker",
                        lambda oid, pid, steer=None: scheduled.update(oid=oid, pid=pid, steer=steer))

    response = client.post(
        f"/profiles/{profile['id']}/opportunities/{opp_id}/draft-full",
        json={"steer": "focus on data governance"},
    )

    assert response.status_code == 202
    assert response.json()["status"] == "drafting"
    assert scheduled["oid"] == opp_id
    assert scheduled["steer"] == "focus on data governance"

    body = client.get(f"/profiles/{profile['id']}/opportunities/{opp_id}").json()
    assert body["package"]["status"] == "drafting"


def test_the_owner_can_edit_every_section():
    """The whole point of a draft is that it gets changed."""
    profile = _profile()
    opp_id = _opportunity(profile["id"])

    edited = client.put(
        f"/profiles/{profile['id']}/opportunities/{opp_id}/package",
        json={"sections": [
            {"title": "Policy Problem Statement", "body": "My own words entirely."},
            {"title": "Research Objectives", "body": "Also mine."},
        ]},
    ).json()

    assert edited["package"]["status"] == "edited"
    assert edited["package"]["sections"][0]["body"] == "My own words entirely."

    history = client.get(
        f"/profiles/{profile['id']}/opportunities/{opp_id}/history"
    ).json()
    assert any(e["kind"] == "edited" for e in history)


def test_exporting_before_drafting_says_so_rather_than_sending_an_empty_file():
    profile = _profile()
    opp_id = _opportunity(profile["id"])

    response = client.get(f"/profiles/{profile['id']}/opportunities/{opp_id}/document")

    assert response.status_code == 409
    assert "draft" in response.json()["detail"]


def test_the_exported_document_is_a_real_docx_with_the_calls_sections():
    profile = _profile()
    opp_id = _opportunity(profile["id"])
    client.put(
        f"/profiles/{profile['id']}/opportunities/{opp_id}/package",
        json={"sections": [{"title": "Policy Problem Statement", "body": "Zimbabwe needs AI policy."}]},
    )

    response = client.get(f"/profiles/{profile['id']}/opportunities/{opp_id}/document")

    assert response.status_code == 200
    # a .docx is a zip - this is the magic number, so it really is one
    assert response.content[:2] == b"PK"
    assert "attachment" in response.headers["content-disposition"]

    from io import BytesIO

    from docx import Document

    doc = Document(BytesIO(response.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Policy Problem Statement" in text
    assert "Zimbabwe needs AI policy." in text


def test_the_exported_document_obeys_the_format_rules_the_call_stated():
    """Times New Roman 12 at 1.5 spacing is grounds for rejection before a
    word is read - applied, not printed as advice."""
    profile = _profile()
    opp_id = _opportunity(profile["id"])

    session = db_module.SessionLocal()
    row = session.get(models_db.StoredOpportunity, opp_id)
    row.package = {
        "spec": {"format_rules": {"font": "Times New Roman", "font_size": 12,
                                  "line_spacing": "1.5"}},
        "sections": [{"title": "Objectives", "body": "Text."}],
    }
    session.commit()
    session.close()

    response = client.get(f"/profiles/{profile['id']}/opportunities/{opp_id}/document")

    from io import BytesIO

    from docx import Document
    from docx.shared import Pt

    doc = Document(BytesIO(response.content))
    normal = doc.styles["Normal"]
    assert normal.font.name == "Times New Roman"
    assert normal.font.size == Pt(12)
    assert normal.paragraph_format.line_spacing == 1.5


def test_another_account_cannot_draft_or_export_yours(outbox):
    profile = _profile()
    opp_id = _opportunity(profile["id"])
    client.cookies.clear()
    sign_up(client, email="other@example.com", outbox=outbox)

    base = f"/profiles/{profile['id']}/opportunities/{opp_id}"
    assert client.post(f"{base}/draft-full", json={}).status_code == 404
    assert client.put(f"{base}/package", json={"sections": []}).status_code == 404
    assert client.get(f"{base}/document").status_code == 404
