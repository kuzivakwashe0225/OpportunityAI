"""MinIO-backed document storage for account profiles.

See SOLUTION_DEFINITION.md §14.2. The client is injectable (same pattern as
search.py's Tavily client): every function here takes a client object rather
than constructing one internally, so tests exercise the real key-construction
and call logic against a lightweight in-memory fake, never a real MinIO
server. make_client() constructs the real one for actual use.
"""

from __future__ import annotations

from datetime import timedelta
from io import BytesIO
from typing import Protocol

DEFAULT_BUCKET = "opportunity-agent-documents"


class ObjectStoreClient(Protocol):
    """The subset of minio.Minio's interface these functions rely on."""

    def bucket_exists(self, bucket_name: str) -> bool: ...
    def make_bucket(self, bucket_name: str) -> None: ...
    def put_object(self, bucket_name: str, object_name: str, data, length: int, content_type: str): ...
    def get_object(self, bucket_name: str, object_name: str): ...
    def presigned_get_object(self, bucket_name: str, object_name: str, expires: timedelta): ...
    def remove_object(self, bucket_name: str, object_name: str) -> None: ...


def make_client(*, endpoint: str, access_key: str, secret_key: str, secure: bool = False):
    from minio import Minio

    return Minio(endpoint, access_key=access_key, secret_key=secret_key, secure=secure)


def ensure_bucket(client: ObjectStoreClient, bucket: str = DEFAULT_BUCKET) -> None:
    if not client.bucket_exists(bucket):
        client.make_bucket(bucket)


def object_key(account_id: str, profile_id: str, document_id: str, filename: str) -> str:
    safe_name = filename.replace("/", "_").replace("\\", "_")
    return f"{account_id}/{profile_id}/{document_id}/{safe_name}"


def upload_document(
    client: ObjectStoreClient,
    *,
    account_id: str,
    profile_id: str,
    document_id: str,
    filename: str,
    content: bytes,
    content_type: str,
    bucket: str = DEFAULT_BUCKET,
) -> str:
    key = object_key(account_id, profile_id, document_id, filename)
    client.put_object(bucket, key, BytesIO(content), length=len(content), content_type=content_type)
    return key


def download_document(client: ObjectStoreClient, key: str, *, bucket: str = DEFAULT_BUCKET) -> bytes:
    response = client.get_object(bucket, key)
    try:
        return response.read()
    finally:
        response.close()
        response.release_conn()


def presigned_url(
    client: ObjectStoreClient,
    key: str,
    *,
    bucket: str = DEFAULT_BUCKET,
    expires: timedelta = timedelta(minutes=15),
) -> str:
    return client.presigned_get_object(bucket, key, expires=expires)


def delete_document(client: ObjectStoreClient, key: str, *, bucket: str = DEFAULT_BUCKET) -> None:
    client.remove_object(bucket, key)


def build_application_docx(
    *,
    title: str,
    sections: list[dict],
    format_rules: dict | None = None,
    applicant: str = "",
    cover_letter: str = "",
) -> bytes:
    """The drafted application as a .docx, obeying the call's own format rules.

    A call that specifies Times New Roman 12 at 1.5 spacing is stating
    grounds for rejection before a word is read, so the rules it stated are
    applied here rather than printed as instructions for the owner to follow
    by hand. Anything the call did not state keeps Word's default - guessing
    a font for a call that never named one would be inventing a requirement.

    Returns bytes rather than writing a file: this is streamed straight to
    the browser and never needs to exist on disk.
    """
    from io import BytesIO

    from docx import Document
    from docx.shared import Pt

    rules = format_rules or {}
    document = Document()

    style = document.styles["Normal"]
    font_name = rules.get("font")
    if font_name:
        style.font.name = str(font_name)
    font_size = rules.get("font_size")
    if font_size:
        try:
            style.font.size = Pt(float(font_size))
        except (TypeError, ValueError):
            pass

    spacing = rules.get("line_spacing")
    if spacing:
        try:
            # Calls write this as "1.5" far more often than as "double".
            style.paragraph_format.line_spacing = float(str(spacing).strip())
        except (TypeError, ValueError):
            pass

    # The letter goes first and on its own page, the way it would be posted.
    # A covering letter stapled into the middle of a proposal is not a
    # covering letter.
    if cover_letter.strip():
        document.add_heading("Covering letter", level=1)
        for para in cover_letter.strip().split("\n\n"):
            cleaned = para.strip()
            if cleaned:
                document.add_paragraph(cleaned)
        document.add_page_break()

    document.add_heading(title, level=1)
    if applicant:
        document.add_paragraph(applicant)

    # Answers to a form's questions are a different kind of thing from the
    # prose sections of a proposal, and running them together produces a
    # document that is neither. They are grouped under their own heading, as
    # question-and-answer, which is the shape the form itself has.
    prose = [s for s in sections if s.get("kind") != "form_field"]
    form_answers = [s for s in sections if s.get("kind") == "form_field"]

    for section in prose:
        heading = str(section.get("title") or "").strip()
        if heading:
            document.add_heading(heading, level=2)
        body = str(section.get("body") or "").strip()
        if body:
            for para in body.split("\n\n"):
                cleaned = para.strip()
                if cleaned:
                    document.add_paragraph(cleaned)
        else:
            # An empty section is left visible on purpose - a gap the owner
            # can see and fill beats a document that quietly omits a section
            # the call requires.
            document.add_paragraph("[This section still needs to be written.]")

    if form_answers:
        document.add_page_break()
        document.add_heading("Application form answers", level=1)
        document.add_paragraph(
            "These answer the questions on the form attached to the call. "
            "Check each one against the form itself before submitting."
        )
        for answer in form_answers:
            question = str(answer.get("title") or "").strip()
            if question:
                document.add_heading(question, level=3)
            body = str(answer.get("body") or "").strip()
            document.add_paragraph(body or "[Still to answer.]")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()

# Enough for a CV, a transcript or a certificate - the documents an applicant
# actually has - without letting a 200-page prospectus fill a small model's
# whole context window. Measured against real CVs: two or three pages of text
# is 4-6k characters, so this holds several times over.
MAX_EXTRACTED_CHARS = 20_000


def readable_text(content: bytes, content_type: str) -> str:
    """The document's words, trimmed to something a prompt can carry.

    Returns "" rather than raising for a file we cannot read - a scanned PDF
    with no text layer is a normal thing for someone to upload, and it must
    not fail their upload or their draft.
    """
    from .document_text import extract_text

    try:
        text = extract_text(content, content_type) or ""
    except Exception:
        return ""
    text = text.strip()
    if len(text) <= MAX_EXTRACTED_CHARS:
        return text
    # Keep the head: a CV puts the qualifications and recent roles first, and
    # the tail is usually referees and hobbies.
    return text[:MAX_EXTRACTED_CHARS]
